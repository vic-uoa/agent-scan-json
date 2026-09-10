from __future__ import annotations

from pathlib import Path
import json
import sys

from docx import Document


def inspect(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append({
                "index": index,
                "style": paragraph.style.name if paragraph.style else "",
                "text": text,
            })

    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    sections = []
    for index, section in enumerate(doc.sections):
        sections.append({
            "index": index,
            "page_width": section.page_width,
            "page_height": section.page_height,
            "top_margin": section.top_margin,
            "bottom_margin": section.bottom_margin,
            "left_margin": section.left_margin,
            "right_margin": section.right_margin,
        })

    return {
        "path": str(path),
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": sections,
        "inline_shapes": len(doc.inline_shapes),
        "styles": [style.name for style in doc.styles if style.type == 1],
    }


if __name__ == "__main__":
    args = sys.argv[1:]
    out_path = None
    if "--out" in args:
        out_index = args.index("--out")
        out_path = Path(args[out_index + 1])
        del args[out_index:out_index + 2]
    payload = [inspect(Path(item)) for item in args]
    rendered = json.dumps(payload, ensure_ascii=False, indent=2)
    if out_path:
        out_path.write_text(rendered, encoding="utf-8")
    else:
        print(rendered)

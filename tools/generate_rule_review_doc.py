from __future__ import annotations

from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCUMENT_SKILL_SCRIPTS = Path(
    r"C:\Users\vic\.cache\codex-runtimes\codex-primary-runtime\plugins\openai-primary-runtime\plugins\documents\skills\documents\scripts"
)
if str(DOCUMENT_SKILL_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(DOCUMENT_SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


ROOT = Path(__file__).resolve().parents[1]
RULES_PATH = ROOT / "skills" / "agent-json-workflow-security-scan" / "rules" / "core-rules.yml"
OUTPUT_PATH = ROOT / "deliverables" / "agent-workflow-json-规则审查与判定细则.docx"

CATEGORY_NAMES = {
    "FLOW": "FLOW｜工作流图与跨节点攻击链",
    "IN": "IN｜输入契约与直接注入面",
    "LLM": "LLM｜模型、Prompt 与输出控制",
    "TOOL": "TOOL｜工具、代码与外部能力",
    "OUT": "OUT｜输出、披露与渲染边界",
    "KB": "KB｜知识检索、ACL 与来源治理",
}


# Each rule has a deliberately minimal scenario.  They describe the scanner's
# static evidence rather than executable payloads, so the document can be used
# as a review guide without becoming a testing playbook for real side effects.
DETAILS: dict[str, tuple[str, str, str]] = {
    "FLOW-001": ("检查 nodes/edges 是否为数组、节点 ID 是否唯一、边两端是否存在。结构错误直接确认。", "nodes: [{id:'a'}, {id:'a'}] 或 edge: a → missing。→ CONFIRMED。", "保持：这是确定性 DSL 完整性错误。"),
    "FLOW-002": ("节点类型未映射，或已知节点含未纳入字段契约的字段时，停止猜测其运行语义。", "data.type: 'NEW_FUTURE_NODE'。→ COVERAGE_GAP，需补平台契约。", "保持：未知字段不靠关键词升级为漏洞。"),
    "FLOW-003": ("引用必须有唯一生产者；必填引用未绑定可确认，歧义/未解析引用作为覆盖缺口。", "MODEL.prompt: '${orderId}'，且无任何 outputName=orderId。→ COVERAGE_GAP。", "保持：区分配置错误与导出信息不足。"),
    "FLOW-004": ("从 HEAD 到副作用/高影响能力必须存在控制路径；动作门须有结构正确、具授权/权限/角色/策略语义且成功分支专属通向动作的 JUDGE。", "HEAD → JUDGE(标题“审批”，条件 route='go') → POST 付款。→ FLOW-004；但 role='admin' 的权限校验分支可作为动作门。", "已复核：不固定为 approved=true，也不接受无授权语义的路由条件。"),
    "FLOW-005": ("知识库或外部工具内容经数据路径进入 MODEL/AGENT，再到副作用工具时形成间接注入链。", "KB.docs → MODEL.prompt → MCP.POST。→ PROBABLE。", "保持：需要连续数据路径，普通只读 RAG 不升级。"),
    "FLOW-006": ("同一入口同时有经过动作门和绕过动作门到达副作用能力的控制路径。", "HEAD → JUDGE(userRole='admin') → POST，且 HEAD → POST。→ CONFIRMED/PROBABLE。", "已复核：动作门采用 FLOW-004 的结构与授权语义证明。"),
    "FLOW-007": ("LOOP 没有最大次数、总时限或停止条件；包含模型/副作用时是运行时覆盖缺口，普通循环仅加固建议。", "LOOP 内含 MODEL，未给 loopCount/maxIterations/timeout。→ COVERAGE_GAP。", "保持：按资源后果分级。"),
    "FLOW-008": ("从具有敏感语义的字段到 external 工具的纯数据路径；敏感性来自字段/嵌套字段/输出名而非任意配置文字。", "HEAD.paramList: email → API_PLUGIN.body。→ PROBABLE。", "已优化：递归识别嵌套字段与 outputName，减少漏报。"),
    "FLOW-009": ("模型输出经数据路径被工具消费：明确 text/String 且到副作用工具时确认；未导出严格 Schema 时仅形成覆盖缺口。", "MODEL.outputFormat='json' 但无 output 字段 → TOOL.body.command。→ COVERAGE_GAP；outputFormat='text' 才可确认自由文本链。", "已复核：不把 JSON 导出缺失误判为运行时自由文本。"),
    "FLOW-010": ("两个不同副作用节点可串联，且整条路径未见幂等、补偿、熔断或失败关闭。", "POST 创建 → POST 扣款，路径无 idempotency/compensation。→ PROBABLE。", "保持：需要实际可达的串联路径。"),
    "FLOW-011": ("WORKFLOW 子工作流调用或 Agent 子工作流工具缺少被调输入输出、委派身份、失败语义。", "WORKFLOW.workflowKey='billing-flow'，未导出被调契约。→ COVERAGE_GAP。", "已优化：区分直接子工作流与 Agent 调用的报告措辞。"),
    "FLOW-012": ("入口可达的 Agent 注册副作用工具，但未导出目标锁定、允许目标或停止边界。", "HEAD → AGENT(tools: POST)，无 goal_lock/kill_switch/max_steps。→ PROBABLE。", "保持：只在 Agent 具有副作用能力时适用。"),
    "FLOW-013": ("非结构节点不从入口可达，或无法到达输出；不把其跨节点链视为可执行。", "孤立 MODEL，无 HEAD→MODEL 或 MODEL→OUTPUT 边。→ OBSERVED/LOW。", "保持：是图卫生信息，不冒充可利用漏洞。"),
    "FLOW-014": ("JUDGE 的 handleId 必须唯一且有对应边；sourceIndex/sourceHandle 必须能映射到导出句柄或显式 ELSE。", "conditionList handleId=[0]，边 sourceHandle='judge-14'。→ CONFIRMED。", "已优化：移除固定 14 作为魔法 ELSE。"),
    "FLOW-015": ("全局 enableHistory 或节点 context=session/history/memory 启用时，检查隔离、TTL、脱敏、容量字段。", "MODEL.context='session'，无 sessionIsolation/ttl/redaction。→ COVERAGE_GAP。", "已优化：覆盖节点级会话上下文，不只看全局开关。"),
    "FLOW-016": ("循环子图必须恰有一个 LOOP_START 与 LOOP_OUTPUT，且两者之间可达。", "LOOP.nodes 仅有 LOOP_START，无 LOOP_OUTPUT。→ CONFIRMED。", "保持：确定性子图契约。"),
    "FLOW-017": ("生产者与消费者显式声明的引用类型不兼容；进入机器消费/副作用后提升严重度。", "MODEL 输出 Array<String> → TEMPLATE.result 声明 String。→ CONFIRMED。", "保持：只比较已声明类型，ANY/未知不臆断。"),
    "IN-001": ("HEAD 必须导出 paramList 和字段 type；Object 还需 sub/children 子字段 Schema。", "paramList: [{name:'profile', type:'Object'}]。→ COVERAGE_GAP（对象键不可验证）。", "已优化：补足标题中“对象 Schema”的实际检查。"),
    "IN-002": ("String/Array 要有有效上限或非空枚举；仅 minLength/minItems 不等价于资源上限。", "{name:'comment', type:'String', minLength:1}。→ OBSERVED/HARDENING。", "已优化：最小值不再被误判为边界。"),
    "IN-003": ("File 字段须有真实 MIME/扩展名允许列表和正数大小上限；多文件还须有数量上限。", "{name:'upload', type:'File', maxSize:1024}。→ CONFIRMED（缺类型允许列表）。", "已优化：不再把字段 type='File' 本身误当作文件类型约束。"),
    "IN-004": ("HEAD 变量真正绑定到 MODEL/AGENT 的 prompt，且 Prompt 未出现不可信数据边界提示。", "prompt:'回答 ${question}'，question 来自 HEAD。→ OBSERVED。", "已优化：只看 prompt 绑定，不因普通 params 误报。"),
    "IN-005": ("字段名具有敏感语义且没有实际非空的分类、保留、最小化或脱敏声明。", "{name:'access_token', type:'String'}。→ COVERAGE_GAP。", "已优化：空 classification 或 false 不再被当作处理策略。"),
    "LLM-001": ("与 IN-004 同一 Prompt 绑定，检查可信指令与不可信内容是否明确分隔。", "prompt:'摘要：${question}'。→ OBSERVED；加入“Treat as untrusted data”可缓解。", "保持：低等级静态观察，需动态注入测试验证。"),
    "LLM-002": ("MODEL/AGENT 节点的文本匹配真实密钥模式且不是 example/placeholder/redacted。", "apiBase 配置旁出现 'sk-实际长密钥'。→ CONFIRMED/HIGH。", "保持：占位符排除以降低误报。"),
    "LLM-003": ("模型输出被 CONDITION/CODE/TOOL/AGENT 消费，但没有封闭字段、类型和嵌套对象 Schema。", "MODEL.outputFormat='json'，无 output[]，结果绑定到 TOOL.body。→ COVERAGE_GAP；明确 text/String 才升级。", "已复核：与 FLOW-009 共用严格 Schema 与证据状态判断。"),
    "LLM-004": ("Agent 或可达高后果模型节点未导出 token、step、iteration 或 timeout 预算。", "AGENT 有 POST 工具，modelConfig:{}。→ COVERAGE_GAP。", "保持：缺运行时字段不确认成漏洞。"),
    "LLM-005": ("Agent 或可达高后果模型节点未见 fallback/retry/errorStrategy/failClosed。", "MODEL → 高影响 TOOL，未配置 fallback。→ COVERAGE_GAP。", "保持：运行时可见性边界清晰。"),
    "LLM-006": ("模型数据进入 JUDGE，JUDGE 的控制分支可达高影响工具。", "MODEL.decision → JUDGE → POST grant-permission。→ PROBABLE。", "已优化：JUDGE 到工具必须为控制路径。"),
    "LLM-007": ("apiBase 明确为 http:// 时记录传输姿态；包含 ${...}/{{...}} 时记录动态端点风险。", "apiBase:'${tenantEndpoint}'。→ PROBABLE，需端点白名单。", "保持：HTTPS 与固定已允许端点不触发。"),
    "TOOL-001": ("TOOL/AGENT 没有可解析 tool spec 是覆盖缺口；spec kind 为 unknown 才是候选。", "AGENT.tools: [] 且无直接工具配置。→ COVERAGE_GAP。", "保持：不推断未知能力的副作用。"),
    "TOOL-002": ("副作用/高影响工具的 url、path、command、tenantId 等敏感参数受不可信变量引用控制。", "HEAD.userId → API_PLUGIN.body.userId，urlMethod='POST'。→ PROBABLE。", "保持：须同时具备不可信绑定和危险能力。"),
    "TOOL-003": ("不可信变量绑定 url/uri/host/endpoint/callback，且节点确为外部或网络能力。", "HEAD.url → API_PLUGIN.query.url。→ PROBABLE/SSRF。", "保持：固定 URL 不适用。"),
    "TOOL-004": ("Python AST 发现 eval/exec/import、进程、网络或文件 IO 原语；动态执行为 CRITICAL。", "code:'eval(params[\"expr\"])'。→ CONFIRMED/CRITICAL。", "保持：只解析代码调用 AST，不扫注释文字。"),
    "TOOL-005": ("外部工具与身份/对象授权相关，但 DSL 未导出 auth/IAM/对象授权事实。", "POST MCP 工具无 auth_declared。→ COVERAGE_GAP。", "保持：运行时 IAM 缺失不能静态确认越权。"),
    "TOOL-006": ("工具输出继续到模型、副作用、高信任输出时，至少一项 spec 未导出 output Schema；CODE 同理。", "外部 GET 输出 → MODEL，spec.output 缺失。→ COVERAGE_GAP。", "保持：仅在下游消费使 Schema 相关时报警。"),
    "TOOL-007": ("副作用工具没有 timeout/retry/idempotency/compensation/circuit breaker 任一可见配置。", "POST 付款无 timeout、retry、idempotency。→ COVERAGE_GAP。", "保持：缺平台策略按覆盖缺口展示。"),
    "TOOL-008": ("可执行且相关的外部工具/MCP/子工作流没有版本，或有版本无签名/摘要。", "MCP version:null，且路径到 MODEL。→ COVERAGE_GAP。", "保持：仅读且无影响的能力不机械触发。"),
    "TOOL-009": ("external TOOL 输出通过数据路径进入 MODEL/AGENT 上下文。", "API_PLUGIN.content → MODEL.prompt。→ OBSERVED/LOW。", "保持：到副作用时由 FLOW-005 升级。"),
    "TOOL-010": ("Agent 可以自主选择副作用工具而 DSL 未证明不可绕过的动作授权策略。", "AGENT.tools 包含 POST。→ PROBABLE。", "保持：仅 Agent + 副作用能力适用。"),
    "TOOL-011": ("Python 无法 AST 解析或代码语言不在已适配范围，危险原语覆盖不完整。", "language:'javascript'。→ COVERAGE_GAP。", "保持：明确说明扫描能力边界。"),
    "TOOL-012": ("工具节点含疑似真实认证材料，或认证参数标为 Agent 可见。", "header.Authorization='Bearer 真实长 token'。→ CONFIRMED；canAgentSee=true → GAP。", "保持：真实值与运行时注入不可见分开报告。"),
    "TOOL-013": ("Python return 推断类型与 CODE.outputType 不同；高后果下游提高严重度。", "return {'ok':true}，outputType:'String'。→ CONFIRMED。", "保持：只对 Python 的可解析 return 做推断。"),
    "OUT-001": ("最终输出没有结构化声明：机器消费/自动决策为风险，人工展示为低等级加固。", "TEMPLATE.outputType:'String' 且 machineConsumed:true。→ PROBABLE/RISK。", "保持：按受众与后果区分。"),
    "OUT-002": ("敏感语义数据有数据路径到最终输出，且输出没有 audience/redaction/masking/accessPolicy。", "HEAD.email → TEMPLATE.result。→ COVERAGE_GAP。", "保持：不能仅由 DSL 断言实际泄露。"),
    "OUT-003": ("INTERMEDIATE_OUTPUT 本身是额外披露边界，未导出受众与脱敏策略。", "INTERMEDIATE_OUTPUT(historyMsgStr:'debug')。→ COVERAGE_GAP。", "保持：中间消息默认需确认受众。"),
    "OUT-004": ("输出明确为 HTML/Markdown，或启用 renderAsHtml/renderLinks，却未见渲染净化策略。", "outputFormat:'html'。→ COVERAGE_GAP。", "已优化：isStreamingOutput=true 的纯文本不再触发。"),
    "OUT-005": ("知识结果可达明确要求引用的输出，但 KB serviceConfig.whetherKnlSource=false。", "KB → TEMPLATE(requiresCitations:true)，来源关闭。→ CONFIRMED。", "保持：没有引用要求时仅 KB-004 加固。"),
    "OUT-006": ("任意声明 JSON/Object 的产出节点缺少字段、字段名/类型重复或对象嵌套 Schema。", "MODEL.outputType:'Object'，output:[]。→ COVERAGE_GAP；重复字段或缺 type → CONFIRMED。", "已复核：缺失导出与显式格式错误分开报告。"),
    "KB-001": ("knowledgeCode 含 ${...} 或 {{...}}，数据集范围可由表达式控制。", "knowledgeCode:'docs-${tenant}'。→ CONFIRMED/HIGH。", "已优化：不再要求动态表达式位于字符串开头。"),
    "KB-002": ("动态数据集，或敏感多数据集，没有可从 DSL 验证的租户 ACL/业务过滤。", "knowledgeCode:'docs-${tenant}'，无 ACL 事实。→ COVERAGE_GAP。", "已优化：运行时 ACL 缺失从 PROBABLE 降为覆盖缺口。"),
    "KB-003": ("知识内容直接绑定 MODEL.prompt，或 Agent 同时绑定知识库与工具；副作用链由 FLOW-005 升级。", "KB.docs → MODEL.prompt。→ OBSERVED/LOW。", "保持：普通 RAG 不是高风险计数。"),
    "KB-004": ("知识可达输出，来源传播关闭，但输出没有明确引用要求。", "whetherKnlSource:false，KB → TEMPLATE。→ OBSERVED/HARDENING。", "保持：避免把可选引用当作漏洞。"),
    "KB-005": ("高后果或敏感知识路径的正数检索量和 0~1 的相关性阈值未完整导出。", "客户 KB → MODEL → POST，recallCount:-1，sortScore:true。→ COVERAGE_GAP。", "已优化：无效数值不再被误认作边界，KB-006 同时确认配置错误。"),
    "KB-006": ("没有固定 knowledgeCode，或检索数量非正整数、score 不在 0~1。", "reCallNum:'zero' 或 sortScore:1.5。→ CONFIRMED。", "保持：值域错误是确定性 DSL 事实。"),
}


def set_east_asia_font(run: Any, name: str) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell: Any, top: int = 70, start: int = 80, bottom: int = 70, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cant_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    tr_pr.append(marker)


def set_cell_width(cell: Any, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_text(cell: Any, text: str, *, bold: bool = False, size: float = 7.7, color: str | None = None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_east_asia_font(run, "Microsoft YaHei")
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.size = Pt(8.5)
    set_east_asia_font(run, "Microsoft YaHei")


def add_section_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(31, 78, 121)
    set_east_asia_font(run, "Microsoft YaHei")


def add_rule_table(document: Document, rules: list[dict[str, Any]]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    table.style = "Table Grid"
    widths = [0.78, 1.62, 4.23, 3.77]
    headers = ["ID", "规则 / 等级", "静态判定逻辑", "最小示例与审查结论"]
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, width, text in zip(header.cells, widths, headers):
        set_cell_width(cell, width)
        shade(cell, "1F4E79")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        add_text(cell, text, bold=True, size=8, color="FFFFFF")

    for index, rule in enumerate(rules):
        rule_id = str(rule["id"])
        logic, example, adjustment = DETAILS[rule_id]
        row = table.add_row()
        set_cant_split(row)
        fill = "F5F9FC" if index % 2 == 0 else "FFFFFF"
        cells = row.cells
        values = [rule_id, f"{rule['title']}\n{rule['severity']}｜{rule['detectability']}", logic, f"{example}\n审查：{adjustment}"]
        for cell, width, value in zip(cells, widths, values):
            set_cell_width(cell, width)
            shade(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            add_text(cell, value, bold=(cell is cells[0]), size=7.5)
    apply_table_geometry(
        table,
        column_widths_from_weights(widths, section_content_width_dxa(document.sections[0])),
        cell_margins_dxa={"top": 70, "start": 80, "bottom": 70, "end": 80},
    )


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("页码 ")
    set_east_asia_font(run, "Microsoft YaHei")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build_document() -> None:
    rules: list[dict[str, str]] = []
    rule_line = re.compile(r"^\s*-\s+\{id:\s*([^,]+),\s*title:\s*([^,]+),\s*severity:\s*([^,]+),\s*detectability:\s*([^,]+),")
    for line in RULES_PATH.read_text(encoding="utf-8").splitlines():
        match = rule_line.match(line)
        if not match:
            continue
        rule_id, title, severity, detectability = (value.strip().strip("'\"") for value in match.groups())
        rules.append({"id": rule_id, "title": title, "severity": severity, "detectability": detectability})
    expected = {str(rule["id"]) for rule in rules}
    if expected != set(DETAILS):
        missing = sorted(expected - set(DETAILS))
        unknown = sorted(set(DETAILS) - expected)
        raise ValueError(f"Rule detail mismatch: missing={missing}, unknown={unknown}")

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.42)
    section.right_margin = Inches(0.42)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if "Compact Note" not in styles:
        style = styles.add_style("Compact Note", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(8.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("AGENT-WORKFLOW-JSON  /  SECURITY RULE REFERENCE")
    header_run.bold = True
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor(31, 78, 121)
    set_east_asia_font(header_run, "Microsoft YaHei")
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("agent-workflow-json 安全规则审查与判定细则")
    run.bold = True
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor(31, 78, 121)
    set_east_asia_font(run, "Microsoft YaHei")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    subrun = subtitle.add_run(f"面向 agent-workflow-json 工作流 DSL ｜审查日期 {date.today().isoformat()} ｜共 {len(rules)} 条规则")
    subrun.font.size = Pt(9)
    subrun.font.color.rgb = RGBColor(89, 89, 89)
    set_east_asia_font(subrun, "Microsoft YaHei")

    add_section_title(document, "阅读说明")
    for item in (
        "本表是静态扫描判定参考。CONFIRMED 表示 DSL 内有确定性错误、危险原语或明确 Text/String 控制链；PROBABLE 表示存在可达攻击链；COVERAGE_GAP 表示需要平台运行时事实，不应当计入已确认漏洞。",
        "示例是最小化 DSL 片段，用于说明规则何时适用；它们不包含可执行攻击载荷。实际报告仍须结合节点路径、JSON Pointer 证据与动态测试计划审阅。",
        "本轮逐条对照 core-rules.yml、json-dsl-bindings.yml、rule-applicability.yml 和 engine.py；优化后的回归集覆盖正例、反例、覆盖缺口与图路径边界。",
    ):
        add_bullet(document, item)

    add_section_title(document, "本轮已落地的判定优化")
    change_table = document.add_table(rows=1, cols=3)
    change_table.autofit = False
    change_widths = [2.1, 4.15, 4.15]
    for cell, width, text in zip(change_table.rows[0].cells, change_widths, ("优化主题", "原有边界", "现在的可验证判定")):
        set_cell_width(cell, width)
        shade(cell, "1F4E79")
        set_cell_margins(cell)
        add_text(cell, text, bold=True, size=8, color="FFFFFF")
    changes = [
        ("动作授权门", "仅凭审批/授权名称即可排除路径", "必须有结构正确、带授权/角色/策略等语义的成功条件分支，且相关分支专属通向动作。"),
        ("结构化输出", "json/object 声明即可压制自由文本控制规则", "要求字段名、类型与嵌套对象 Schema；Schema 缺失是覆盖缺口，明确 text/String 才确认自由文本链。"),
        ("条件句柄", "固定数值 14 被默认为 ELSE", "仅接受导出 handleId、可验证序号映射或显式 ELSE。"),
        ("输入与文件", "min 视为上限；File 类型名视为 MIME 限制", "要求正数上限/非空枚举，以及真实文件类型允许列表和大小上限。"),
        ("富文本输出", "流式文本一律视为富文本", "只在 HTML、Markdown、链接或 HTML 渲染明确开启时检查净化。"),
        ("知识库运行时事实", "未见 ACL/检索配置时产生可能风险", "ACL 与运行时检索策略按 COVERAGE_GAP；无效数值仍确定性确认。"),
    ]
    for index, values in enumerate(changes):
        row = change_table.add_row()
        set_cant_split(row)
        fill = "F5F9FC" if index % 2 == 0 else "FFFFFF"
        for cell, width, value in zip(row.cells, change_widths, values):
            set_cell_width(cell, width)
            shade(cell, fill)
            set_cell_margins(cell)
            add_text(cell, value, size=7.7)
    apply_table_geometry(
        change_table,
        column_widths_from_weights(change_widths, section_content_width_dxa(section)),
        cell_margins_dxa={"top": 70, "start": 80, "bottom": 70, "end": 80},
    )

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for rule in rules:
        grouped[str(rule["id"]).split("-", 1)[0]].append(rule)
    for prefix in ("FLOW", "IN", "LLM", "TOOL", "OUT", "KB"):
        document.add_page_break()
        add_section_title(document, CATEGORY_NAMES[prefix])
        intro = document.add_paragraph(style="Compact Note")
        intro.paragraph_format.space_after = Pt(4)
        run = intro.add_run(f"{len(grouped[prefix])} 条规则。每一行给出适用前提、报告证据状态与一个最小判定例。")
        run.font.size = Pt(8.5)
        set_east_asia_font(run, "Microsoft YaHei")
        add_rule_table(document, grouped[prefix])

    document.add_page_break()
    add_section_title(document, "审查结论与使用边界")
    for item in (
        "保留 54 条规则 ID 与既有标准映射；本轮优化聚焦判定精度、覆盖范围与证据状态，不改变规则目录的兼容性。",
        "扫描报告应将风险、姿态、加固建议和覆盖缺口分列展示；覆盖缺口必须在补足平台运行时证据后才能升级为漏洞结论。",
        "高后果工作流建议先修复 CONFIRMED/PROBABLE 链，再补齐 IAM、ACL、渲染器净化、供应链签名和资源预算等 COVERAGE_GAP。",
    ):
        add_bullet(document, item)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
